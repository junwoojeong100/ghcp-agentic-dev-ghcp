"""Check the rendered deck and produce contact sheets for visual inspection."""
from pathlib import Path
import json
import re
import fitz
from PIL import Image, ImageDraw
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
PRES = ROOT / "presentation"
RENDERED = PRES / "rendered"
RENDERED.mkdir(exist_ok=True)
pptx = Presentation(PRES / "GitHub-Copilot-Agentic-Development.pptx")
pdf = fitz.open(PRES / "GitHub-Copilot-Agentic-Development.pdf")
normalize = lambda value: re.sub(r"\s+", "", value)


def visual_text(page, shape):
    rect = fitz.Rect(shape.left / 12700, shape.top / 12700,
                     (shape.left + shape.width) / 12700, (shape.top + shape.height) / 12700)
    rows = []
    for block in page.get_text("rawdict")["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                for char in span["chars"]:
                    bbox = fitz.Rect(char["bbox"])
                    center = fitz.Point((bbox.x0 + bbox.x1) / 2, (bbox.y0 + bbox.y1) / 2)
                    if not center in rect:
                        continue
                    baseline = char["origin"][1]
                    row = next((row for row in rows if abs(row["baseline"] - baseline) < max(2, span["size"] * .3)), None)
                    if row is None:
                        row = {"baseline": baseline, "chars": []}
                        rows.append(row)
                    row["chars"].append((bbox.x0, char["c"]))
    return "\n".join("".join(char for _, char in sorted(row["chars"], key=lambda item: item[0])) for row in sorted(rows, key=lambda row: row["baseline"]))

report = {"pptxSlides": len(pptx.slides), "pdfPages": len(pdf), "missingText": [], "outsidePage": [],
          "notes": [], "introductionChecks": [], "structureErrors": []}
if len(pptx.slides) != len(pdf):
    raise RuntimeError("PowerPoint and PDF page counts differ.")
structure = json.loads((PRES / "slide-audit.json").read_text())
intro_numbers = structure["introduction"]["slide_numbers"]
intro_requirements = [
    ["GitHub Copilot", "AI", "코드"],
    ["IDE", "CLI", "GitHub.com", "Copilot app", "GitHub Mobile"],
    ["Ask", "Plan", "Agent", "Edit", "Interactive", "Autopilot", "권한"],
    ["AGENT MODE / CLI", "CLOUD AGENT", "COPILOT CODE REVIEW", "소개"],
    ["Prompt files", "Custom agents", "Subagents", "Fleet", "Agent skills", "MCP", "Hooks", "미시연"],
]
if intro_numbers != list(range(2, len(intro_requirements) + 2)):
    report["structureErrors"].append("The five introduction slides must immediately follow the cover.")
for number, required in zip(intro_numbers, intro_requirements):
    slide = pptx.slides[number - 1]
    actual = normalize("\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)).casefold()
    missing = [token for token in required if normalize(token).casefold() not in actual]
    sources_present = "https://docs.github.com/" in slide.notes_slide.notes_text_frame.text
    report["introductionChecks"].append({"slide": number, "missing": missing, "officialSourcesInNotes": sources_present})
video_slides = [
    index for index, slide in enumerate(pptx.slides, 1)
    if any(shape.click_action.hyperlink.address == "../video/Agentic-Development-Demo-KO.mp4" for shape in slide.shapes)
]
if video_slides != [structure["demo_video_slide"]]:
    report["structureErrors"].append("The video handoff does not match the declared slide number.")
if len(pptx.slides) != structure["main_slides"] + structure["technical_appendix_slides"]:
    report["structureErrors"].append("Main and appendix counts do not match the actual deck.")
thumbs = []
for index, (slide, page) in enumerate(zip(pptx.slides, pdf), 1):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text.strip():
            expected = normalize(shape.text)
            if expected not in normalize(visual_text(page, shape)):
                report["missingText"].append({"slide": index, "text": shape.text})
    for block in page.get_text("dict")["blocks"]:
        for row in block.get("lines", []):
            for span in row["spans"]:
                rect = fitz.Rect(span["bbox"])
                if rect.x0 < -1 or rect.y0 < -1 or rect.x1 > page.rect.width + 1 or rect.y1 > page.rect.height + 1:
                    report["outsidePage"].append({"slide": index, "text": span["text"], "box": list(rect)})
    report["notes"].append(bool(slide.notes_slide.notes_text_frame.text.strip()))
    pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
    path = RENDERED / f"slide-{index:02d}.png"
    pixmap.save(path)
    image = Image.open(path)
    image.thumbnail((480, 270))
    thumb = Image.new("RGB", (500, 304), "#dce2ec")
    thumb.paste(image, (10, 10))
    ImageDraw.Draw(thumb).text((13, 284), f"SLIDE {index:02d}", fill="#182236")
    thumbs.append(thumb)
for start in range(0, len(thumbs), 12):
    subset = thumbs[start:start + 12]
    rows = (len(subset) + 3) // 4
    sheet = Image.new("RGB", (2000, rows * 304), "#dce2ec")
    for i, image in enumerate(subset):
        sheet.paste(image, ((i % 4) * 500, (i // 4) * 304))
    sheet.save(RENDERED / f"contact-sheet-{start//12+1}.png")
report["passed"] = (not report["missingText"] and not report["outsidePage"] and all(report["notes"])
                    and not report["structureErrors"]
                    and all(not check["missing"] and check["officialSourcesInNotes"] for check in report["introductionChecks"]))
(ROOT / "evidence/deck-check.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
print(json.dumps({k: v for k, v in report.items() if k != "notes"}, ensure_ascii=False, indent=2))
if not report["passed"]:
    raise RuntimeError("Inspect the rendered presentation; text, introduction content, or slide structure checks failed.")
