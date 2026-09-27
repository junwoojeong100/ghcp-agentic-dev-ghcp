"""Build the Korean Copilot introduction and CXO deck with verified CLI evidence."""
from pathlib import Path
import json
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "presentation"
ASSETS = OUT / "assets"
STATE = json.loads((ROOT / "demo/cli-reference/.demo/state.json").read_text())
BROWSER = json.loads((ROOT / "evidence/browser-check.json").read_text())
EDIT = json.loads((ROOT / "video/terminal/edit.json").read_text())
CAPTURE = json.loads((ROOT / f"video/terminal/{EDIT['runId']}/capture.json").read_text())
if not (STATE["verification"]["passed"] and BROWSER["passed"] and CAPTURE["complete"]):
    raise RuntimeError("Build only from the actual, completed CLI recording and verified source.")
if BROWSER["sourceDigest"] != STATE["verification"]["sourceDigest"] or STATE["id"] != CAPTURE["runId"]:
    raise RuntimeError("Slides, browser evidence, and CLI recording must describe the same run.")

W, H = 13.333333, 7.5
FONT = "Apple SD Gothic Neo"
C = {
    "ink": "182236", "dark": "101727", "panel": "1A2439", "paper": "F7F8FC",
    "white": "FFFFFF", "muted": "78859A", "line": "DCE2EC", "purple": "A783D7",
    "lilac": "D6C1F0", "mint": "64D5AF", "amber": "F0C083", "red": "D36C78",
    "pale": "EEE7F7", "green": "267F62", "lighttext": "B1BED3",
}
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = "GitHub Copilot: 제품 소개부터 실제 에이전트 실행까지"
prs.core_properties.subject = "CXO 브리핑: Copilot 개요 · 서피스 · 작업 모드 · 에이전트 활용 · 실제 CLI 데모"
prs.core_properties.author = "GitHub Copilot presenter kit"
prs.core_properties.keywords = "GitHub Copilot, surfaces, modes, CLI, cloud agent, custom agents, CXO"


def rgb(value):
    return RGBColor.from_string(C.get(value, value))


def shape(slide, x, y, w, h, fill, line_color=None, rounded=False):
    obj = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(fill)
    if line_color:
        obj.line.color.rgb = rgb(line_color)
        obj.line.width = Pt(.8)
    else:
        obj.line.fill.background()
    if rounded:
        obj.adjustments[0] = .12
    effect = obj._element.find(f"{qn('p:style')}/{qn('a:effectRef')}")
    if effect is not None:
        effect.set("idx", "0")
    return obj


def text(slide, content, x, y, w, h, size=18, color="ink", bold=False,
         align=PP_ALIGN.LEFT, font=FONT, line_spacing=1.13):
    obj = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = obj.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for index, value in enumerate(content.split("\n")):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.text = value
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(4 if size >= 16 else 2)
        for run in paragraph.runs:
            run.font.name = "Arial" if font == FONT else font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
            properties = run._r.get_or_add_rPr()
            east_asian = etree.SubElement(properties, qn("a:ea"))
            east_asian.set("typeface", FONT)
    return obj


def pill(slide, content, x, y, w, color="purple", dark=False):
    shape(slide, x, y, w, .34, "panel" if dark else "pale", rounded=True)
    text(slide, content, x + .08, y + .045, w - .16, .25, 10, color, True, PP_ALIGN.CENTER)


def line(slide, x1, y1, x2, y2, color="line"):
    obj = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    obj.line.color.rgb = rgb(color)
    obj.line.width = Pt(.7)
    effect = obj._element.find(f"{qn('p:style')}/{qn('a:effectRef')}")
    if effect is not None:
        effect.set("idx", "0")
    return obj


def arrow(slide, x, y, color="purple", w=.35):
    obj = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(.3))
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(color)
    obj.line.fill.background()


def new_slide(eyebrow, title, subtitle="", dark=False, title_size=32):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("dark" if dark else "paper")
    text(slide, eyebrow.upper(), .59, .34, 12.0, .26, 10, "lilac" if dark else "purple", True)
    text(slide, title, .59, .91, 12.1, 1.65 if title_size >= 40 else 1.23,
         title_size, "white" if dark else "ink", True, line_spacing=1.09)
    if subtitle:
        text(slide, subtitle, .61, 2.13, 12.0, .45, 15, "lighttext" if dark else "muted")
    line(slide, .59, 7.05, 12.73, 7.05, "panel" if dark else "line")
    text(slide, "GitHub Copilot  /  CXO Briefing. Real execution. Human control.", .61, 7.16, 10.8, .19, 8.5, "muted")
    text(slide, f"{len(prs.slides):02d}", 12.08, 7.12, .63, .26, 11, "muted", align=PP_ALIGN.RIGHT)
    return slide


def notes(slide, body, sources=""):
    slide.notes_slide.notes_text_frame.text = (
        body + "\n\n표현 경계: 실제 Copilot CLI PTY 녹화와 합성 주문 앱을 사용했습니다. "
        "영상의 승인 키 입력은 자동 리허설이며 실제 사람·조직의 승인이 아닙니다. "
        "라이브 session 명령은 발표자 입력을 기다립니다. 실제 결제·PR·배포 없음. "
        "편집 영상 길이는 개발 시간이나 생산성 개선율이 아닙니다.\n"
        + (f"\n출처/근거:\n{sources}" if sources else "")
    )


def picture(slide, path, x, y, w, h):
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    shape(slide, x, y, w, h, "white" if path.name in {"before-result.png", "after-result.png"} else "dark", rounded=True)
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                                   width=Inches(pw), height=Inches(ph))


def crop(source, dest, box):
    Image.open(ASSETS / source).crop(box).save(ASSETS / dest)
    return ASSETS / dest


def cards(slide, items, dark=False, y=2.78, height=3.22):
    for i, (label, title, body, color) in enumerate(items):
        x = .62 + i * 4.15
        shape(slide, x, y, 3.88, height, "panel" if dark else "white",
              None if dark else "line", True)
        pill(slide, label, x + .24, y + .24, 3.39, color, dark)
        text(slide, title, x + .25, y + .97, 3.36, .82, 24,
             "white" if dark else "ink", True)
        text(slide, body, x + .25, y + 2.01, 3.36, 1.02, 17,
             "lighttext" if dark else "muted")


before = crop("01-before.png", "before-result.png", (647, 342, 1330, 877))
after = crop("02-after.png", "after-result.png", (647, 342, 1330, 900))
repository = next(s for s in STATE["verification"]["suites"] if s["name"] == "repository")
acceptance = next(s for s in STATE["verification"]["suites"] if s["name"] == "acceptance")
baseline = next(s for s in STATE["baselineTests"]["suites"] if s["name"] == "acceptance")
CLI = "demo/cli-reference/.demo/"
DOCS = "https://docs.github.com/en/copilot"
INTRO_SOURCES = {
    "overview": f"{DOCS}/get-started/about-github-copilot",
    "ide": f"{DOCS}/how-tos/chat-with-copilot/chat-in-ide",
    "matrix": f"{DOCS}/reference/copilot-feature-matrix",
    "cli": f"{DOCS}/concepts/agents/copilot-cli/about-copilot-cli",
    "autopilot": f"{DOCS}/concepts/agents/copilot-cli/autopilot",
    "app": f"{DOCS}/get-started/quickstart-copilot-app",
    "mobile": f"{DOCS}/how-tos/copilot-on-github/chat-with-copilot/chat-in-mobile",
    "cloud": f"{DOCS}/concepts/agents/cloud-agent/about-cloud-agent",
    "review": f"{DOCS}/concepts/agents/code-review",
    "customization": f"{DOCS}/reference/customization-cheat-sheet",
    "fleet": f"{DOCS}/how-tos/copilot-cli/use-copilot-cli/speed-up-task-completion",
}

# 01
s = new_slide("GITHUB COPILOT / CXO BRIEFING", "GitHub Copilot\n개발 생산성의 다음 단계.", dark=True, title_size=42)
text(s, "개발자의 시간을 더 가치 있는 일에.\n속도·품질·통제를 함께 설계합니다.", .66, 2.74, 6.13, 1.16, 24, "lilac", True)
text(s, "반복 작업은 Copilot에게\n범위와 판단은 개발자에게\n고객 문제는 확인 가능한 결과로", .67, 4.22, 6.08, 1.52, 22, "white")
pill(s, "ACTUAL CLI · CUSTOM AGENTS · HUMAN IN THE LOOP", .67, 6.28, 6.06, "mint", True)
shape(s, 7.38, 2.71, 5.28, 3.72, "panel", rounded=True)
text(s, "고객의 같은 요청", 7.75, 3.05, 4.51, .48, 20, "lighttext", align=PP_ALIGN.CENTER)
text(s, "2회 → 주문 1건", 7.72, 3.83, 4.61, .88, 34, "mint", True, PP_ALIGN.CENTER)
text(s, "실제 CLI 실행 영상으로 확인합니다.", 7.79, 5.18, 4.43, .69, 19, "white", align=PP_ALIGN.CENTER)
notes(s, "CXO 오프닝: 먼저 GitHub Copilot이 무엇인지, 어디서 사용하고 어떤 방식으로 "
      "작업을 맡길 수 있는지 짧게 소개합니다. 이어 팀의 실행 여력을 어떻게 확장할지 "
      "보겠습니다. 반복 작업을 AI에 맡기고 개발자는 중요한 결정에 집중하는 방식입니다. "
      "생산성, 가치 전달 속도, 품질과 통제의 잠재 가치를 실제 고객 문제 해결로 보여 줍니다. "
      "같은 인원으로 몇 퍼센트 더 많이 처리했다는 성과는 측정하지 않았으므로 주장하지 않습니다.",
      "video/Agentic-Development-Demo-KO.mp4")

# 02
s = new_slide("01 / BUSINESS VALUE", "Copilot에 기대하는 것은\n더 많은 코드가 아니라, 더 큰 실행력입니다.")
cards(s, [
    ("생산성", "핵심 개발에 집중", "반복 탐색·수정은 AI에\n개발자는 핵심 판단에 집중", "purple"),
    ("가치 전달 속도", "요청에서 실행까지", "계획·변경·검토를 연결\n수작업 인계를 줄이는 구조", "green"),
    ("품질과 통제", "맡기되, 확인 가능", "승인 지점과 근거를 남겨\n변경의 위험을 확인", "purple"),
])
text(s, "관찰할 것: 실제 작업 위임과 승인 흐름  |  측정할 것: 리드타임·재작업·검토 부담의 변화",
     .68, 6.39, 12, .44, 17, "ink", True)
notes(s, "맥락, 실행, 통제의 세 축으로 설명합니다. 컨텍스트 복사나 반복 작업을 줄일 수 "
      "있는 구조지만 이 자료는 대조 실험을 수행하지 않았습니다. 수치 대신 실제 파일 읽기, "
      "편집 승인, 코드 변경, 검토의 증거를 뒤에서 보여 줍니다.",
      f"{CLI}logs/\n{DOCS}/concepts/agents/about-copilot-cli")

# 03
s = new_slide("02 / OPERATING MODEL", "사람을 빼는 자동화가 아니라,\n팀의 일하는 방식을 확장합니다.", dark=True)
for y, label, items, color in [
    (2.78, "COPILOT에게 위임", ["코드 읽기", "해결 계획", "파일 수정", "근거 검토"], "lilac"),
    (4.82, "개발자가 결정", ["완료 기준", "허용 범위", "편집 승인", "최종 판단"], "mint"),
]:
    text(s, label, .7, y, 11.9, .37, 17, color, True)
    for i, item in enumerate(items):
        x = .67 + i * 3.16
        shape(s, x, y + .59, 2.57, .85, "panel", rounded=True)
        text(s, item, x + .18, y + .81, 2.22, .44, 21, "white", True, PP_ALIGN.CENTER)
        if i < 3:
            arrow(s, x + 2.74, y + .86, color, .24)
notes(s, "사람이 없어지는 흐름이 아니라 일과 결정의 분담입니다. 업무 기준과 승인 책임은 "
      "사람에게 남고 에이전트는 읽기, 계획, 실제 변경, 검토를 수행합니다. 테스트 실행은 "
      "모델이 아니라 로컬 Node.js 도구가 수행한다는 점도 구분합니다.", "demo/scripts/workflow.mjs")

# 04
s = new_slide("03 / FROM CUSTOMER IMPACT TO EXECUTION", "고객 문제 한 건으로,\n개발 조직의 실행력을 확인합니다.")
shape(s, .65, 2.77, 4.10, 3.71, "white", "line", True)
pill(s, "합성 고객 문의", .93, 3.04, 3.48, "red")
text(s, "“응답이 없어서\n다시 보냈을 뿐인데,\n주문이 두 건이에요.”", .95, 3.79, 3.42, 1.55, 24, "ink", True)
text(s, "고객 신뢰와 문의 대응 부담", .96, 5.75, 3.43, .63, 16, "muted")
picture(s, before, 5.10, 2.77, 7.57, 3.71)
notes(s, "변경 전 앱에서 실제 동시 요청 두 개가 서로 다른 주문 두 건을 만듭니다. "
      "실제 고객사 장애나 결제 사례가 아닌 합성 문제입니다. 업무 결과부터 보여 주고, "
      "그다음 실제 CLI로 해결하는 흐름으로 연결합니다.", "evidence/browser-check.json")

# 05
s = new_slide("04 / HUMAN-OWNED ACCEPTANCE CRITERIA", "중복은 막고,\n정상적인 새 주문은 유지합니다.")
cards(s, [
    ("같은 요청 재전송", "2회 → 1건", "동시에 도착해도\n같은 주문번호로 응답", "purple"),
    ("별도의 새 주문", "새 주문 허용", "다른 요청 표식이면\n정상적으로 추가 접수", "green"),
    ("같은 표식 · 다른 내용", "충돌 거부", "상품·수량이 달라지면\n기존 주문을 바꾸지 않음", "red"),
])
text(s, "사람이 정한 경계: 코드·테스트 4개 파일만  |  결제·DB·인증·의존성·배포 변경 없음",
     .66, 6.40, 12.05, .42, 16, "ink", True)
notes(s, "요청서의 기준은 구현 에이전트가 바꿀 수 없습니다. 버튼을 비활성화하는 것만으로 "
      "문제를 숨기지 않고 서버에서 재전송을 구분합니다. 허용 파일, 금지 작업, 완료 기준을 "
      "명확히 한 다음 범위를 승인합니다.", "demo/starter/REQUEST.md\ndemo/acceptance/duplicate-order.test.mjs")

# 06
s = new_slide("05 / WORK DELEGATION WITH ACCOUNTABILITY", "AI는 역할을 나눠 일하고,\n사람은 범위와 결과를 통제합니다.", dark=True)
stages = [
    ("계획 담당 AI", "요청·코드 읽기\n계획 제안", "demo-planner", "lilac"),
    ("구현 담당 AI", "승인된 파일의\n코드·테스트 수정", "demo-implementer", "lilac"),
    ("실제 테스트", "고정된 기준 실행\n결과·해시 기록", "Node.js / local helper", "mint"),
    ("검토 담당 AI", "변경·테스트 읽기\n위험·누락 검토", "demo-reviewer", "lilac"),
]
for i, (title, body, command, color) in enumerate(stages):
    x = .64 + i * 3.17
    shape(s, x, 2.77, 2.64, 2.32, "panel", rounded=True)
    text(s, title, x + .19, 3.01, 2.26, .44, 20, color, True)
    text(s, body, x + .19, 3.65, 2.27, .77, 17, "white")
    text(s, command, x + .19, 4.65, 2.27, .30, 10.2, "lighttext", font="Arial")
    if i < 3:
        arrow(s, x + 2.79, 3.74, "muted", .23)
pill(s, "구현 전: 범위 승인 + 네이티브 파일 편집 승인", .67, 5.49, 7.16, "amber", True)
pill(s, "검토 후: 사람의 최종 판단", 8.03, 5.49, 4.62, "amber", True)
text(s, "역할별 실제 CLI 세션을 로컬 파일로 인계합니다. 자율 A2A 통신 프로토콜 구현은 아닙니다.",
     .67, 6.43, 12, .38, 14, "lighttext")
notes(s, "세 역할은 실제 .github/agents 프로필입니다. 각 역할을 별도의 Copilot CLI "
      "세션으로 실행하고 로컬 컨트롤러가 산출물을 인계합니다. 이 로컬 순서 제어를 "
      "Copilot의 내장 오케스트레이터라고 소개하지 않습니다.",
      f"{CLI}logs/\ndemo/scripts/workflow.mjs\n{DOCS}/reference/custom-agents-configuration")

# 07
s = new_slide("06 / REAL PRODUCT DEMONSTRATION", "AI가 실제로 일하고,\n사람이 통제하는 장면을 보겠습니다.",
              "코드를 읽지 않아도 됩니다. 고객 결과, 승인 지점, 검토 근거만 확인해 주세요.")
picture(s, ASSETS / "cli-00-start.png", .65, 2.80, 12.02, 3.57)
obj = text(s, "5분 영상 재생  ▶", .70, 6.49, 3.6, .42, 20, "purple", True)
obj.click_action.hyperlink.address = "../video/Agentic-Development-Demo-KO.mp4"
text(s, "원본 PTY 녹화 · 승인 입력은 자동 리허설 · 대기 편집/가속", 4.28, 6.54, 8.22, .35, 13, "muted")
notes(s, "여기에서 MP4를 재생합니다. 대시보드로 CLI를 흉내 낸 화면이 아니라 실제 "
      "Copilot 프로세스의 터미널 출력입니다. 원본 캐스트와 입력 기록도 제공합니다. "
      "영상의 승인 키 입력은 자동 리허설임을 숨기지 않습니다. 라이브 session 명령은 "
      "실제 발표자의 입력을 기다립니다.", "video/README.ko.md\nvideo/terminal/edit.json")

# 08
s = new_slide("APPENDIX A / CLI EVIDENCE: CONTEXT", "코드를 읽는 단계부터\n개발 작업을 위임합니다.")
picture(s, ASSETS / "cli-01-plan.png", .64, 2.82, 8.54, 3.68)
for i, (title, body) in enumerate([
    ("요청 + 실제 저장소", "서버·화면·테스트를 함께 확인"),
    ("변경 영향 연결", "재전송·정상 주문·오류 상태 고려"),
    ("실행 가능한 계획", "허용된 네 파일로 범위 구체화"),
]):
    y = 2.98 + i * 1.16
    text(s, title, 9.51, y, 3.12, .42, 19, "purple", True)
    text(s, body, 9.51, y + .55, 3.08, .57, 15, "muted")
notes(s, "실제 계획 역할의 터미널 화면입니다. CLI 하단의 demo-planner와 파일 읽기 결과를 "
      "보여 줍니다. 맥락을 찾고 변경 영향을 연결하는 반복 작업을 맡길 수 있다는 이점을 "
      "말하되 모든 저장소 이해가 항상 정확하다고 주장하지 않습니다.", f"{CLI}01-plan.md\n{CLI}logs/planner.events.jsonl")

# 09
s = new_slide("08 / CONTROLLED AI ADOPTION", "작업은 AI에 맡겨도,\n승인 권한은 개발자에게 남깁니다.", dark=True)
picture(s, ASSETS / "cli-03-permission.png", .65, 2.80, 8.45, 3.76)
text(s, "① 범위 승인", 9.43, 2.95, 3.19, .44, 23, "amber", True)
text(s, "로컬 게이트\n명시적 문구 없으면 구현 차단", 9.45, 3.54, 3.12, .79, 17, "white")
text(s, "② 편집 승인", 9.43, 4.58, 3.19, .44, 23, "mint", True)
text(s, "Copilot의 실제 권한 UI\n변경을 보고 ‘이번만 허용’", 9.45, 5.17, 3.12, .90, 17, "white")
notes(s, "두 승인 계층을 엄격히 구분합니다. 범위 문구를 받는 것은 키트의 로컬 게이트이고, "
      "스크린샷의 파일 편집 메뉴는 Copilot CLI의 실제 제품 UI입니다. 모든 파일 작업을 "
      "미리 허용하지 않았고 첫 번째 Yes를 선택했습니다. 역할별 도구 제한은 실제 기능이지만 "
      "조직 인증이나 완전한 보안 격리를 대신하지 않습니다.",
      "demo/scripts/cli-session.mjs\nvideo/terminal/" + EDIT["runId"] + "/actions.jsonl")

# 10
s = new_slide("APPENDIX B / CLI EVIDENCE: EXECUTION", "설명으로 끝내지 않고,\n저장소의 코드와 테스트를 바꿉니다.")
picture(s, ASSETS / "cli-04-implementation.png", .64, 2.81, 8.56, 3.76)
text(s, "실제로 바뀐 영역", 9.50, 2.97, 3.15, .45, 21, "purple", True)
text(s, "서버 중복 방지\n요청 상태·재사용 표시\n고객 화면 안내\n회귀 테스트", 9.51, 3.76, 3.09, 1.92, 20, "ink")
text(s, "범위 밖 변경은 별도로 검사", 9.51, 6.01, 3.10, .52, 15, "muted")
notes(s, "실제 구현 역할이 파일을 편집한 기록입니다. 파일 네 개의 diff와 로그를 제공하며 "
      "모의 성공 문구로 대체하지 않습니다. 일관된 여러 파일 변경을 위임할 수 있는 장점을 "
      "강조합니다. 실제 결제·DB·인증·배포를 변경한 것은 아닙니다.", f"{CLI}02-diff.patch\n{CLI}02-implementation.md")

# 11
s = new_slide("APPENDIX C / CLI EVIDENCE: QUALITY", "테스트 통과로 끝내지 않고,\n브라우저 실패까지 반영합니다.", dark=True)
for x, title, value, detail in [
    (.67, "사람이 고정한 업무 기준", f"{acceptance['pass']} / {acceptance['tests']}", f"변경 전 {baseline['fail']}개 실패 → 수정 후 통과"),
    (4.83, "저장소 회귀 테스트", f"{repository['pass']} / {repository['tests']}", "기존 동작과 추가 경계 조건 확인"),
]:
    shape(s, x, 2.83, 3.84, 3.20, "panel", rounded=True)
    text(s, title, x + .27, 3.12, 3.3, .44, 18, "lighttext", True)
    text(s, value, x + .27, 3.89, 3.3, .93, 46, "mint", True)
    text(s, detail, x + .28, 5.18, 3.26, .59, 15, "white")
text(s, "첫 서버 테스트 통과\n→ 브라우저 실패 발견\n→ CLI 재수정·재확인", 9.32, 3.46, 3.28, 1.89, 20, "white", True)
text(s, "테스트는 로컬 실행 도구가 수행합니다. AI의 자기평가나 결함 부재의 보증이 아닙니다.",
     .69, 6.48, 12, .35, 14, "lighttext")
notes(s, "숫자는 현재 녹화한 실행의 최종 테스트 기록에서 읽어 생성합니다. 첫 구현은 "
      "서버 테스트를 통과했고 검토도 브라우저 미확인을 명시한 채 검토 준비 권고를 했습니다. "
      "이후 브라우저에서 오류 코드 안내 누락과 실패 상태의 정상 배너를 관찰했습니다. "
      "그 실제 기록을 구현 역할에 전달해 재수정하고 다시 확인했습니다. 검토 에이전트가 "
      "그 결함을 먼저 찾았다고 말하지 않습니다. 업무 기준 테스트는 수정 범위 밖에 있고 "
      "테스트 개수는 생산성 수치가 아닙니다.",
      f"{CLI}00-baseline.tap\n{CLI}03-tests.tap\n{CLI}03-verification.json")

# 12
s = new_slide("07 / CUSTOMER-VISIBLE VALUE", "AI의 가치는,\n고객에게 보이는 결과로 확인합니다.")
pill(s, "BEFORE · 요청 2회 / 주문 2건", .66, 2.65, 5.79, "red")
pill(s, "AFTER · 요청 2회 / 주문 1건", 6.90, 2.65, 5.77, "green")
picture(s, before, .66, 3.16, 5.79, 3.32)
picture(s, after, 6.90, 3.16, 5.77, 3.32)
notes(s, "이 화면은 녹화된 CLI가 변경한 소스의 실행 결과입니다. 이전 데모의 성공 "
      "화면을 끼워 넣지 않았고 소스 해시를 연결해 확인했습니다. 별도 새 주문이 "
      "누적 두 번째 주문을 만드는 것도 브라우저에서 확인했습니다.", "evidence/browser-check.json\nevidence/video-recording.json")

# 13
s = new_slide("09 / QUALITY AND ACCOUNTABILITY", "품질은 확인 가능하게.\n출시는 책임 있게.")
picture(s, ASSETS / "cli-06-review.png", .65, 2.81, 8.49, 3.76)
text(s, "Copilot 검토", 9.48, 2.99, 3.13, .45, 22, "purple", True)
text(s, "실제 변경·테스트 확인\n누락·위험·미확인 조건 정리", 9.49, 3.59, 3.13, .95, 17, "ink")
text(s, "개발자 최종 판단", 9.48, 4.94, 3.17, .46, 22, "green", True)
text(s, "검토 권고 ≠ 출시 승인\n자동 PR·병합·배포 없음", 9.49, 5.55, 3.13, .90, 17, "ink")
notes(s, "실제 읽기 전용 검토 역할의 결과입니다. 다른 역할로 검토한다고 결함이 "
      "모두 발견된다고 보장하지 않습니다. 단일 프로세스 메모리 구현의 운영 한계와 "
      "사람의 최종 출시 판단을 남기는 것이 중요합니다.", f"{CLI}04-review.md\n{CLI}state.json")

# 14
s = new_slide("10 / CXO TAKEAWAY", "Copilot은 개발 도구를 넘어,\n팀의 실행력을 확장하는 선택입니다.", dark=True)
for i, (title, body) in enumerate([
    ("생산성 · 반복 작업 위임", "탐색·구현 부담을 맡기고 핵심 개발에 집중"),
    ("속도 · 작업 흐름 연결", "요청부터 실제 변경·검토까지 이어지는 실행"),
    ("품질 · 근거 중심 판단", "실패와 재작업까지 기록하고 결과를 확인"),
    ("통제 · 책임 있는 확장", "승인 지점과 허용 범위를 정해 AI 활용"),
]):
    x = .68 + (i % 2) * 6.18
    y = 2.79 + (i // 2) * 1.74
    shape(s, x, y, 5.89, 1.49, "panel", rounded=True)
    text(s, title, x + .27, y + .28, 5.29, .48, 22, "mint" if i > 1 else "lilac", True)
    text(s, body, x + .28, y + .94, 5.23, .35, 16, "white")
notes(s, "효과를 과장하는 수치 없이 네 가지 관찰 가능한 강점을 정리합니다. 코드를 "
      "추천하는 도구에서 승인된 개발 작업을 수행하는 파트너로 확장되는 가치를 설명합니다. "
      "조직의 실제 시간·품질·검토 부담 변화는 별도 파일럿으로 측정해야 합니다.")

# 15
s = new_slide("11 / INVESTMENT DECISION", "한 팀, 한 업무로 시작하고,\n확장은 성과를 보고 결정합니다.")
cards(s, [
    ("속도", "가치 전달 리드타임", "요청부터 검토까지 걸린 시간\n유사한 작업의 전후 비교", "purple"),
    ("품질", "재작업과 검토 부담", "요구 누락·회귀 결함\n사람의 검토 시간 함께 확인", "green"),
    ("경제성", "회수한 시간의 가치", "라이선스·사용량·도입 비용과\n확보한 개발 여력을 함께 평가", "purple"),
])
text(s, "활용도는 채택 지표, ROI는 업무 성과와 비용의 문제입니다. 절감률은 파일럿에서 확인합니다.",
     .68, 6.42, 12.0, .42, 17, "ink", True)
notes(s, "CXO의 다음 결정은 전사 일괄 확대가 아니라 한 팀·한 업무·한 책임자를 정한 "
      "파일럿 승인입니다. 기준선과 완료 기준을 합의하고 유사 복잡도의 작업을 비교합니다. "
      "GitHub Copilot usage metrics는 채택과 활동, PR 흐름의 참고 데이터가 될 수 있지만 "
      "사용량 증가만으로 ROI나 인과 효과를 입증하지는 않습니다. 회수된 시간의 추정 가치와 "
      "실제 현금 지출 절감도 구분하고 교육·검토·도입 비용을 포함합니다. 이 키트에서 조직 "
      "metrics 대시보드를 연결하거나 조직 성과를 측정한 것은 아닙니다.",
      f"{DOCS}/concepts/billing-and-usage/copilot-usage-metrics/copilot-metrics")

# 16
s = new_slide("APPENDIX D / PRODUCT VS DEMO HELPER", "제품 기능과 데모 보조 도구를 구분합니다.")
for x, title, body, dark in [
    (.66, "실제 GitHub Copilot CLI", "네이티브 터미널 UI\n역할별 custom agent 프로필\n파일 읽기·편집·변경 검토\n실제 파일 편집 권한 요청", False),
    (6.90, "이 키트의 로컬 보조 도구", "역할 순서 제어·파일 인계\n명시적 범위 승인 문구\n변경 범위·해시 검사와 테스트 실행\n원본 PTY 녹화·편집·재생", True),
]:
    shape(s, x, 2.73, 5.78, 3.74, "dark" if dark else "white", None if dark else "line", True)
    text(s, title, x + .3, 3.03, 5.16, .54, 23, "lilac" if dark else "green", True)
    text(s, body, x + .31, 3.89, 5.13, 2.07, 20, "white" if dark else "ink")
text(s, "A2A 프로토콜·Copilot Studio·GitHub Actions 연결 아님. 로컬 승인은 조직 인증을 대신하지 않음.",
     .69, 6.66, 12.0, .28, 12, "muted")
notes(s, "승인 문구를 받는 로컬 도구를 Copilot 내장 승인 기능과 혼동하지 않습니다. "
      "파일별 실제 권한 요청은 Copilot 제품 화면입니다. 원본 터미널 출력은 xterm.js로 "
      "렌더링하되 모델 응답이나 승인 UI를 새로 만들어 넣지 않았습니다.",
      "demo/scripts/cli-session.mjs\ntools/capture-cli.py\ntools/render-terminal.mjs")

# 17
s = new_slide("APPENDIX E / RUN IT, THEN CHECK THE EVIDENCE", "실연 명령과 원본 근거를 함께 제공합니다.")
text(s, "실제 CLI + 발표자 승인", .67, 2.69, 5.91, .48, 23, "purple", True)
text(s, "cd demo\nnpm run demo -- session live-01", .70, 3.39, 5.70, .93, 17, "ink", font="Menlo")
text(s, "Git·Node.js 22+·Copilot CLI·gh 로그인 필요\n계획 확인 → 승인 문구 → 파일별 Yes\n각 역할 응답 후 /exit로 다음 단계 인계", .71, 4.64, 5.66, 1.34, 17, "muted")
text(s, "원본과 결과", 7.03, 2.69, 5.65, .48, 23, "green", True)
for i, (label, path) in enumerate([
    ("영상", "video/Agentic-Development-Demo-KO.mp4"),
    ("원본 터미널", f"video/terminal/{EDIT['runId']}/session.cast"),
    ("실제 변경·로그", "demo/cli-reference/.demo/"),
    ("발표·복구 가이드", "DEMO-GUIDE.ko.md"),
]):
    y = 3.42 + i * .61
    text(s, label, 7.06, y, 1.51, .35, 13, "muted", True)
    text(s, path, 8.62, y, 4.06, .39, 10.2, "ink", font="Arial")
text(s, "오프라인 결과: cd demo && npm run fallback", 7.07, 6.16, 5.51, .4, 12.5, "ink", font="Arial")
notes(s, "새 실행은 작업 공간을 덮어쓰지 않으므로 매번 새 ID를 사용합니다. "
      "실제 모델 사용에는 계정 정책과 사용량 과금이 적용될 수 있습니다. CLI 또는 "
      "네트워크가 막히면 저장본임을 알리고 MP4 또는 fallback으로 전환합니다.",
      "README.md\nDEMO-GUIDE.ko.md\nvideo/README.ko.md")

# 18
s = new_slide("APPENDIX F / SOURCES & CLAIM BOUNDARIES", "확인한 기능만, 근거와 함께 설명합니다.")
text(s, "공식 GitHub 문서", .68, 2.7, 5.75, .48, 23, "ink", True)
sources = [
    ("GitHub Copilot 제품 개요", INTRO_SOURCES["overview"]),
    ("IDE의 작업 모드·에이전트", INTRO_SOURCES["ide"]),
    ("CLI 기능·권한 경계", INTRO_SOURCES["cli"]),
    ("Cloud agent·비동기 위임", INTRO_SOURCES["cloud"]),
    ("Agents·Skills·MCP·Hooks", INTRO_SOURCES["customization"]),
    ("도입 활용도·업무 흐름 metrics", f"{DOCS}/concepts/billing-and-usage/copilot-usage-metrics/copilot-metrics"),
]
for i, (label, url) in enumerate(sources):
    obj = text(s, f"{i+1:02d}  {label} ↗", .73, 3.35 + i * .40, 5.62, .36, 16, "purple")
    obj.click_action.hyperlink.address = url
text(s, "실행 환경: Copilot CLI 1.0.88\n확인일: 2026-09-26 KST", .73, 6.0, 5.56, .70, 15, "muted")
shape(s, 6.91, 2.72, 5.74, 3.91, "dark", rounded=True)
text(s, "이 데모로 보장하지 않는 것", 7.24, 3.08, 5.10, .58, 22, "amber", True)
text(s, "생산성·매출·비용 개선율\n결함이 전혀 없다는 보장\n재시작·분산 환경·실제 결제 안전성\n인증된 조직 승인·자동 배포",
     7.25, 4.00, 5.06, 1.91, 19, "white")
notes(s, "2026-09-26에 공식 GitHub 문서와 설치된 CLI 도움말을 확인했습니다. "
      "서두는 제품 전반의 개요이고, 이후 CLI 데모는 실제로 실행한 범위만 보여 줍니다. "
      "제품 기능과 키트 보조 기능, 관측과 기대 효과, 리허설 입력과 실제 승인을 구분합니다. "
      "각 서두 슬라이드 노트에는 해당 서피스·모드·기능의 상세 출처가 있습니다.",
      "\n".join(url for _, url in sources) + "\nevidence/VERIFICATION.ko.md")

existing_order = [1, 2, 3, 4, 5, 6, 7, 12, 9, 13, 14, 15, 8, 10, 11, 16, 17, 18]
existing_count = len(prs.slides)

# These introductions are placed immediately after the cover.
s = new_slide("INTRO 01 / WHAT IS GITHUB COPILOT", "GitHub Copilot은,\n소프트웨어 개발을 돕는 AI입니다.",
              "코드 작성·이해·계획·변경·검토를 기존 개발 흐름 안에서 지원합니다.")
cards(s, [
    ("ASSIST", "코드를 제안하고", "입력 중 코드 추천\n질문 답변·코드 설명", "purple"),
    ("CONTEXT", "맥락을 이해하고", "요청·코드·저장소를 참고\n작업에 필요한 정보 탐색", "green"),
    ("AGENT", "작업을 실행합니다", "목표를 받아 여러 단계 수행\n변경과 검증 결과를 검토", "purple"),
])
text(s, "자동완성만이 아니라, 개발 업무를 함께 수행하는 파트너입니다. 최종 판단은 사람에게 남습니다.",
     .67, 6.40, 12.0, .42, 16, "ink", True)
notes(s, "권장 40초. GitHub는 코드와 협업을 관리하는 플랫폼이고 GitHub Copilot은 "
      "그 개발 흐름에서 코드 작성·이해·출시 준비를 돕는 AI 제품입니다. 자동완성과 대화형 "
      "보조뿐 아니라 목표를 받고 파일과 도구를 다루는 에이전트 기능도 제공합니다. "
      "이 제품을 Microsoft 365 Copilot이나 Copilot Studio와 같은 것으로 설명하지 않습니다. "
      "지원 기능은 플랜·클라이언트·조직 정책에 따라 다르며 모든 서피스가 같은 기능을 "
      "제공하는 것은 아닙니다. 이 슬라이드는 제품 소개이지 모든 기능의 실행 증거는 아닙니다.",
      INTRO_SOURCES["overview"])

s = new_slide("INTRO 02 / SURFACES: WHERE YOU USE COPILOT", "필요한 작업을,\n익숙한 접점에서 시작합니다.",
              "서피스(Surface)는 Copilot을 사용하는 화면·도구입니다.")
surfaces = [
    ("IDE", "작성·질문·수정\n열린 프로젝트와\n함께 작업", "VS Code·VS·JetBrains 등"),
    ("CLI", "터미널에서\n계획·코드 수정\n도구 실행", "이번 데모의 접점"),
    ("GitHub.com", "저장소·이슈·PR에서\n질문·작업 위임\n변경 검토", "웹 기반 협업"),
    ("Copilot app", "전용 데스크톱 앱\n프로젝트·에이전트\n세션 관리", "에이전트 작업 공간"),
    ("GitHub Mobile", "이동 중 저장소 탐색\n코드·PR 질문\n검토 지원", "모바일 접점"),
]
for i, (title, body, foot) in enumerate(surfaces):
    x = .65 + i * 2.44
    active = title == "CLI"
    shape(s, x, 2.83, 2.27, 3.19, "dark" if active else "white", None if active else "line", True)
    text(s, f"0{i+1}", x + .19, 3.07, 1.87, .34, 15, "mint" if active else "purple", True)
    text(s, title, x + .19, 3.62, 1.91, .43, 19, "white" if active else "ink", True)
    text(s, body, x + .19, 4.32, 1.89, 1.02, 14, "lighttext" if active else "muted")
    text(s, foot, x + .19, 5.53, 1.89, .33, 10.5, "mint" if active else "muted", True)
text(s, "CLI는 여러 접점 중 하나입니다. 지원 기능은 클라이언트·버전·플랜·조직 정책에 따라 다릅니다.",
     .67, 6.42, 12.0, .38, 14, "ink", True)
notes(s, "권장 45초. 서피스는 사용 접점, 작업 모드는 도움을 받는 방식, 실행 환경은 도구가 "
      "실제로 동작하는 위치입니다. 서로 같은 분류가 아닙니다. IDE에는 VS Code, Visual Studio, "
      "JetBrains 등이 있으며 세부 지원은 feature matrix를 봅니다. GitHub.com에서는 저장소 "
      "대화, Cloud agent 위임, PR 검토 등을 사용할 수 있습니다. Copilot app은 별도 데스크톱 "
      "제품으로 프로젝트와 에이전트 세션을 관리하며 GitHub Desktop과 혼동하지 않습니다. "
      "GitHub Mobile은 저장소·코드·PR 질문과 지원되는 검토 기능을 제공합니다. "
      "앱과 CLI의 조직 정책도 별도일 수 있습니다. 이 데모는 강조한 CLI만 실제 실행했습니다.",
      "\n".join(INTRO_SOURCES[key] for key in ["overview", "matrix", "cli", "app", "mobile", "review"]))

s = new_slide("INTRO 03 / MODES: HOW YOU WORK", "작업에 맞게,\n질문·계획·수정·실행을 선택합니다.")
for i, (mode, title, body, color) in enumerate([
    ("Ask", "질문과 이해", "코드 설명·탐색\n접근 방법 상담", "purple"),
    ("Plan", "구현 전 계획", "요구·범위·단계 정리\n계획을 보고 구현 결정", "green"),
    ("Agent", "목표 기반 실행", "파일 수정·명령 실행\n결과를 보고 반복", "purple"),
    ("Edit*", "선택 범위 수정", "지정한 파일을 중심으로\n변경안을 제안·적용", "muted"),
]):
    x = .65 + i * 3.10
    shape(s, x, 2.70, 2.78, 2.10, "white", "line", True)
    text(s, mode, x + .21, 2.93, 2.35, .43, 23, color, True)
    text(s, title, x + .21, 3.53, 2.35, .38, 18, "ink", True)
    text(s, body, x + .21, 4.08, 2.35, .62, 14, "muted")
text(s, "CLI", .71, 5.19, .87, .43, 22, "purple", True)
for i, (mode, body) in enumerate([
    ("Interactive", "대화하며 실행"),
    ("Plan", "구현 전에 계획"),
    ("Autopilot", "여러 단계를 연속 진행"),
]):
    x = 1.83 + i * 3.66
    shape(s, x, 5.02, 3.44, .96, "pale", rounded=True)
    text(s, mode, x + .17, 5.15, 3.08, .34, 17, "ink", True)
    text(s, body, x + .17, 5.60, 3.08, .25, 12.5, "muted")
text(s, "모드명·제공 범위는 클라이언트별로 다릅니다. Edit*은 지원 IDE·버전에 한정됩니다.",
     .69, 6.20, 12.0, .30, 12.5, "muted")
text(s, "연속 진행과 도구 권한은 별도입니다. 이번 녹화는 Interactive + 수동 편집 승인입니다.",
     .69, 6.57, 12.0, .30, 13, "ink", True)
notes(s, "권장 60초. 위 네 카드는 IDE에서 접하는 대표 작업 방식이며 모든 클라이언트에 "
      "동일한 메뉴가 있다는 뜻은 아닙니다. 현재 GitHub의 VS Code 사용 가이드는 Ask·Plan·Agent를 "
      "소개합니다. Edit은 feature matrix에 나온 지원 IDE·버전에서 선택 파일을 중심으로 "
      "수정하는 방식이므로 별표로 한정했습니다. CLI는 대화형과 프로그램 방식 인터페이스를 "
      "제공하며, 대화형 작업에서 Interactive·Plan·Autopilot을 구분할 수 있습니다. "
      "Autopilot은 완료나 제한에 이를 때까지 이어서 작업하는 모드이지 권한 무제한과 같은 "
      "옵션이 아닙니다. 제한된 권한으로 계속하면 승인이 필요한 작업이 거부될 수 있고, "
      "넓은 권한 부여에는 별도의 위험 검토가 필요합니다. 이 데모는 Autopilot을 사용하지 "
      "않았습니다. 영상의 demo-planner는 사용자 정의 역할이지 내장 Plan 모드를 켠 증거가 아닙니다.",
      "\n".join(INTRO_SOURCES[key] for key in ["ide", "matrix", "cli", "autopilot"]))

s = new_slide("INTRO 04 / AGENT-POWERED WORK", "에이전트는 목표를 받아,\n개발 작업을 여러 단계로 수행합니다.", dark=True)
cards(s, [
    ("AGENT MODE / CLI", "구현·검증을 위임", "소스 탐색·여러 파일 수정\n도구 실행·테스트와 재시도", "lilac"),
    ("CLOUD AGENT", "백그라운드에 위임", "GitHub에서 조사·계획·변경\n필요하면 PR로 전달", "mint"),
    ("COPILOT CODE REVIEW", "변경을 다시 검토", "IDE·PR의 변경을 검토\n문제와 개선안 제안", "lilac"),
], dark=True, height=3.08)
text(s, "IDE의 Agent mode와 Cloud agent는 별개입니다. AI의 검토 의견은 출시 승인이 아닙니다.",
     .69, 6.12, 12.0, .34, 14, "white", True)
text(s, "이번 시연: CLI의 Custom agent + 로컬 테스트. Cloud 위임·PR 리뷰 서비스는 소개만 합니다.",
     .69, 6.58, 12.0, .30, 12.5, "lighttext")
notes(s, "권장 50초. Agent mode나 CLI는 개발 환경에서 파일과 도구를 사용해 목표를 해결할 "
      "수 있습니다. Cloud agent는 GitHub의 별도 실행 환경에서 백그라운드로 조사·계획·변경을 "
      "수행하고 필요할 때 PR을 만드는 서비스입니다. IDE의 Agent mode와 같은 이름의 모드로 "
      "묶지 않습니다. GitHub의 현재 명칭은 Copilot cloud agent이며 이전 coding agent 명칭을 "
      "아는 참석자에게는 같은 제품의 현재 명칭으로 설명합니다. Copilot code review는 "
      "변경 맥락을 모아 문제와 개선안을 제안하는 제품 기능이며 검토를 사람의 승인으로 "
      "해석하지 않습니다. 이번 영상은 세 역할의 CLI 실행을 로컬 파일로 인계한 사례입니다. "
      "테스트는 로컬 컨트롤러가 실행했고 demo-reviewer는 사용자 정의 읽기 전용 역할입니다. "
      "Cloud agent나 GitHub의 PR code review 서비스를 실제로 실행한 데모는 아닙니다.",
      "\n".join(INTRO_SOURCES[key] for key in ["ide", "cli", "cloud", "review"]))

s = new_slide("INTRO 05 / TEAM-READY AGENT FEATURES", "우리 팀의 역할·도구·절차에 맞춰\n에이전트를 확장할 수 있습니다.")
features = [
    ("팀의 기준을 공유", "Instructions / Prompt files", "규칙과 반복 요청을 재사용"),
    ("역할을 전문화", "Custom agents", "담당 역할·지침·도구를 구성"),
    ("작업을 나눠 수행", "Subagents / Fleet", "하위 작업 위임·가능한 부분 병렬화"),
    ("업무 절차를 재사용", "Agent skills", "필요할 때 절차·스크립트·자료 로드"),
    ("외부 도구와 연결", "MCP servers", "이슈·문서·브라우저 등의 도구 연결"),
    ("실행 지점에 통제", "Hooks", "도구 실행 전후 검사·기록"),
]
for i, (title, feature, body) in enumerate(features):
    x = .65 + (i % 3) * 4.15
    y = 2.65 + (i // 3) * 1.68
    shape(s, x, y, 3.86, 1.48, "white", "line", True)
    text(s, title, x + .23, y + .18, 3.37, .41, 20, "ink", True)
    text(s, feature, x + .24, y + .72, 3.36, .27, 12, "purple", True, font="Arial")
    text(s, body, x + .24, y + 1.12, 3.36, .26, 12.5, "muted")
shape(s, .66, 6.17, 12.04, .57, "pale", rounded=True)
text(s, "이번 사용: 저장소 지침·Custom agents·편집 승인  |  그 외 확장 기능은 제품 소개이며 미시연",
     .89, 6.31, 11.57, .30, 13, "ink", True)
notes(s, "권장 60초. Custom instructions는 작업 맥락에 적용할 기준이고 Prompt files는 "
      "반복 요청 템플릿입니다. Prompt files를 CLI 공통 기능이라고 말하지 않으며 지원 "
      "클라이언트를 확인합니다. Custom agents는 역할·도구 구성이며 Subagents는 별도 "
      "맥락에서 하위 작업을 수행하는 실행 단위입니다. CLI의 Fleet는 병렬 가능한 작업을 "
      "subagents에 나누는 기능으로, 에이전트 수가 늘면 항상 빨라진다고 보장하지 않습니다. "
      "Skills는 관련 업무의 절차·스크립트·자료를 필요할 때 로드합니다. MCP는 외부 도구와 "
      "데이터에 연결하고 Hooks는 실행 지점에서 검사·기록·정책 동작을 수행하도록 확장합니다. "
      "지원 범위는 클라이언트·버전·정책에 따라 다르고 지침 자체가 보안 격리를 보장하지는 "
      "않습니다. 이번 촬영은 저장소 지침, 세 Custom agent 프로필, 네이티브 편집 승인을 "
      "사용했습니다. Fleet·Subagents·Skills·MCP·Hooks를 시연한 것으로 설명하지 않습니다. "
      "이제 이 제품 소개를 바탕으로 다음 CXO 가치와 실제 CLI 사례로 연결합니다.",
      "\n".join(INTRO_SOURCES[key] for key in ["customization", "fleet", "ide"]))

introduction_count = len(prs.slides) - existing_count
order = [existing_order[0], *range(existing_count + 1, len(prs.slides) + 1), *existing_order[1:]]
slide_ids = list(prs.slides._sldIdLst)
for number in order:
    prs.slides._sldIdLst.append(slide_ids[number - 1])
for number, slide in enumerate(prs.slides, 1):
    for item in slide.shapes:
        if item.has_text_frame and item.top >= Inches(7.1) and item.left > Inches(12):
            item.text_frame.paragraphs[0].runs[0].text = f"{number:02d}"

OUT.mkdir(exist_ok=True)
target = OUT / "GitHub-Copilot-Agentic-Development.pptx"
prs.save(target)
audit = {
    "slides": len(prs.slides),
    "audience": "CXO customers",
    "main_slides": len(prs.slides) - 6,
    "technical_appendix_slides": 6,
    "introduction": {
        "slide_numbers": list(range(2, introduction_count + 2)),
        "topics": ["Copilot definition", "Surfaces", "Modes", "Agent-powered work", "Agent customization"],
        "sources_verified_on": "2026-09-26",
        "sources": INTRO_SOURCES,
    },
    "demo_video_slide": existing_order.index(7) + introduction_count + 1,
    "all_have_notes": all(bool(slide.notes_slide.notes_text_frame.text.strip()) for slide in prs.slides),
    "source_digest": STATE["verification"]["sourceDigest"],
    "cli_run": CAPTURE["runId"],
    "cast_sha256": CAPTURE["castSha256"],
    "actual_test_counts": {"acceptance": acceptance["tests"], "repository": repository["tests"]},
    "off_slide_shapes": [],
}
for index, slide in enumerate(prs.slides, 1):
    for item in slide.shapes:
        if item.left < 0 or item.top < 0 or item.left + item.width > prs.slide_width + 3 or item.top + item.height > prs.slide_height + 3:
            audit["off_slide_shapes"].append({"slide": index, "name": item.name})
(OUT / "slide-audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n")
if audit["off_slide_shapes"]:
    raise RuntimeError(f"Off-slide content: {audit['off_slide_shapes']}")
print(f"Saved {target} ({len(prs.slides)} editable slides with notes and genuine CLI screenshots)")
